import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm

from catalog_generator import export_catalog, write_catalog_docx


class CatalogTemplateTests(unittest.TestCase):
    def test_table_marker_split_runs_preserves_template(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            template = Document()
            template.sections[0].left_margin = Mm(31)
            template.sections[0].header.paragraphs[0].text = "Marca do template"
            template.sections[0].footer.paragraphs[0].text = "Rodapé original"
            template.add_paragraph("Antes")
            cell = template.add_table(rows=1, cols=1).cell(0, 0)
            cell.paragraphs[0].add_run("{CATA")
            cell.paragraphs[0].add_run("LOGO}")
            template.add_paragraph("Depois")
            source = root / "template.docx"
            template.save(source)
            before = source.read_bytes()
            output = root / "output.docx"
            write_catalog_docx(output, [], template_path=source)
            generated = Document(output)
            self.assertEqual(source.read_bytes(), before)
            self.assertIn("Modelo não identificado", generated.tables[0].cell(0, 0).text)
            self.assertNotIn("{CATALOGO}", generated.tables[0].cell(0, 0).text)
            self.assertEqual(generated.sections[0].header.paragraphs[0].text, "Marca do template")
            self.assertEqual(generated.sections[0].footer.paragraphs[0].text, "Rodapé original")
            self.assertAlmostEqual(generated.sections[0].left_margin.mm, 31, places=1)
            self.assertFalse(generated._element.xpath('.//w:br[@w:type="page"]'))

    def test_inline_or_duplicate_marker_is_not_silently_ignored(self):
        for markers in (("Texto {CATALOGO}",), ("{CATALOGO}", "{CATALOGO}")):
            with self.subTest(markers=markers), tempfile.TemporaryDirectory() as temp:
                source = Path(temp) / "template.docx"
                document = Document()
                for text in markers:
                    document.add_paragraph(text)
                document.save(source)
                with self.assertRaisesRegex(ValueError, "marcador"):
                    write_catalog_docx(Path(temp) / "out.docx", [], source)

    def test_template_pdf_uses_generated_docx_and_never_generic_fallback(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "template.docx"
            document = Document()
            document.add_paragraph("IDENTIDADE ORIGINAL")
            document.save(source)
            def convert(docx, pdf):
                self.assertIn("IDENTIDADE ORIGINAL", [p.text for p in Document(docx).paragraphs])
                pdf.write_bytes(b"%PDF-test-converter")
            exports = export_catalog(root, {}, [], "f" * 32, source, convert)
            self.assertIn("pdf", exports)
            failed = Mock(side_effect=RuntimeError("Word indisponível"))
            exports = export_catalog(root, {}, [], "f" * 32, source, failed)
            self.assertEqual(set(exports), {"docx", "json", "csv", "xlsx"})
            failed.assert_called_once()


if __name__ == "__main__":
    unittest.main()
