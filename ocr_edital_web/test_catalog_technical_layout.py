import tempfile
import unittest
import zipfile
from pathlib import Path

import pdfplumber
from docx import Document
from catalog_generator import export_catalog, normalize_items, write_catalog_docx
from catalog_technical_layout import STANDARD_CLAUSES, SECTIONS, technical_sections, technical_pending_questions


class TechnicalCatalogTests(unittest.TestCase):
    def test_unproven_clauses_are_explicitly_pending_not_asserted(self):
        entry = {"nome": "Cadeira PP", "itens": ["3"], "caracteristicas": ["Assento em polipropileno"]}
        sections = technical_sections(entry, entry["caracteristicas"])
        text = "\n".join(line for _, lines in sections for line in lines)
        self.assertEqual([title for title, _ in sections], [title for title, _ in SECTIONS])
        for _, clause in STANDARD_CLAUSES:
            self.assertIn("Cláusula padrão pendente de confirmação (não é declaração de atendimento): " + clause, text)
        self.assertNotIn("Conformidade declarada na evidência:", text)
        self.assertIn("Item nº: 3", text)
        self.assertTrue(technical_pending_questions(entry))

    def test_exact_evidence_keeps_required_terms_and_negation_does_not_certify(self):
        clauses = [clause for _, clause in STANDARD_CLAUSES]
        entry = {"nome": "Modelo confirmado", "caracteristicas": clauses}
        text = "\n".join(line for _, lines in technical_sections(entry, clauses) for line in lines)
        for clause in clauses:
            self.assertIn(clause, text)
        self.assertNotIn("Cláusula padrão pendente", text)
        denied = ["Não foi confirmada: " + clause for clause in clauses]
        text = "\n".join(line for _, lines in technical_sections(entry, denied) for line in lines)
        self.assertNotIn("Conformidade declarada na evidência:", text)
        self.assertIn("Cláusula padrão pendente", text)

    def test_answers_appear_in_the_correct_fields_without_inventing_values(self):
        entry = {"nome": "Modelo", "respostas_layout": {"Cor": "Azul", "Garantia": "24 meses"}}
        sections = dict(technical_sections(entry, []))
        self.assertIn("Cor: Azul", sections["RESUMO DO ITEM"])
        self.assertIn("Garantia: 24 meses", sections["OBSERVAÇÕES"])

    def test_component_presence_does_not_supply_unreported_dimensions(self):
        entry = {"caracteristicas": ["Pistão classe 4 e base com cinco rodízios."]}
        questions = technical_pending_questions(entry)
        requirements = [question["requisito"] for question in questions]
        self.assertIn("Diâmetro dos rodízios (mm), quando aplicável", requirements)
        self.assertIn("Pistão a gás, quando aplicável", requirements)
        sections = dict(technical_sections({**entry, "nome": "Modelo"}, entry["caracteristicas"]))
        self.assertIn(entry["caracteristicas"][0], sections["MECANISMOS E ACESSÓRIOS"])

    def test_both_exports_have_identity_sections_and_preserve_evidence_boundary(self):
        items = normalize_items([{"numeroItem": 7, "descricao": "Cadeira giratória em tela Mesh com apoio lombar e pistão classe 4. SEGREDO_REQUISITO", "quantidade": 1, "unidade": "UN"}], "offline:test")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            exports = export_catalog(root, {"numero_compra": "45/2026", "processo": "112/2026", "orgao": "Municipio de Teste", "objeto": "OBJETO_NAO_ESPECIFICACAO"}, items, "layout-test")
            document = Document(root / exports["docx"]["filename"])
            word_text = "\n".join(p.text for p in document.paragraphs)
            with pdfplumber.open(root / exports["pdf"]["filename"]) as pdf:
                pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
                self.assertTrue(all(page.width < page.height for page in pdf.pages))
            for text in (word_text, pdf_text):
                for title, _ in SECTIONS:
                    self.assertIn(title, text)
                self.assertIn("45/2026", text)
                self.assertIn("112/2026", text)
                self.assertIn("Municipio de Teste", text)
                self.assertIn("Item nº: 7", text)
                self.assertNotIn("SEGREDO_REQUISITO", text)
                self.assertNotIn("OBJETO_NAO_ESPECIFICACAO", text)

    def test_new_layout_preserves_all_non_body_template_parts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            doc = Document()
            doc.sections[0].header.paragraphs[0].text = "CABECALHO ORIGINAL"
            doc.sections[0].footer.paragraphs[0].text = "RODAPE ORIGINAL"
            doc.add_paragraph("{CATALOGO}")
            template = root / "template.docx"
            output = root / "out.docx"
            doc.save(template)
            entry = {"nome": "Modelo", "familia": "Cadeira", "caracteristicas": [], "dimensoes": []}
            write_catalog_docx(output, [entry], template, {"numero_compra": "45/2026"})
            with zipfile.ZipFile(template) as before, zipfile.ZipFile(output) as after:
                self.assertEqual(before.namelist(), after.namelist())
                for name in before.namelist():
                    if name != "word/document.xml":
                        self.assertEqual(before.read(name), after.read(name), name)
            self.assertIn("ASSENTO E ENCOSTO", "\n".join(p.text for p in Document(output).paragraphs))
