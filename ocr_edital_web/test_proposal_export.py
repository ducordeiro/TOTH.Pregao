import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

import server


class ProposalExportTests(unittest.TestCase):
    def test_preview_metrics_match_export_with_custom_template_margins(self):
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / "template.docx"
            output = Path(directory) / "output.docx"
            doc = Document()
            doc.sections[0].left_margin = Inches(1.5)
            doc.sections[0].right_margin = Inches(1.25)
            doc.styles["Normal"].paragraph_format.left_indent = Pt(20)
            doc.styles["Normal"].paragraph_format.space_before = Pt(15)
            doc.save(template)
            with patch.object(server, "resolve_template", return_value=template):
                metrics = server.docx_structure_response({"template_ref": "test"})["table_metrics"]
            items = [{"item": "2", "quantidade": "32", "unidade": "UND", "descricao": "Cadeira fixa",
                      "marca": "Teste", "valor_unitario": "R$ 234,00", "valor_total": "R$ 7.488,00"}]
            requested = {"item": 6, "quantidade": 6, "unidade": 6, "descricao": 52,
                         "marca": 8, "valor_unitario": 11, "valor_total": 11}
            server.build_docx(items, template, output, proposal_column_widths=requested)
            table = Document(output).tables[-1]
            total = sum(column.width.twips for column in table.columns)
            self.assertEqual(total, metrics["available_width_twips"])
            for column, weight in zip(table.columns, requested.values()):
                self.assertAlmostEqual(column.width.twips / total * 100, weight, delta=0.03)
            for index, row in enumerate(table.rows):
                for cell in row.cells:
                    p = cell.paragraphs[0]
                    self.assertEqual(p.paragraph_format.left_indent.pt, 0)
                    self.assertEqual(p.paragraph_format.space_before.pt, 0)
                    self.assertEqual(p.runs[-1].font.size.pt, metrics["header_font_pt" if index == 0 else "body_font_pt"])
                padding = row.cells[-1]._tc.tcPr.find(qn("w:tcMar")).find(qn("w:start"))
                self.assertEqual(int(padding.get(qn("w:w"))), metrics["horizontal_padding_twips"])

    def context(self):
        return {
            "items": [{"item": "1", "descricao": "Cadeira"}],
            "template_path": Path("original.docx"), "source_name": "teste",
            "responsible_id": "1", "responsible": {}, "commercial_terms": {},
            "mini_box_order": ["box"], "document_block_order": ["box", "table"],
            "mini_box_alignments": {"box": "center"}, "mini_box_contents": {"box": "Titulo"},
            "extra_column": None, "table_layout": {"columns": []},
            "proposal_column_widths": {"descricao": 50},
        }

    def test_word_remains_default_without_conversion(self):
        with tempfile.TemporaryDirectory() as directory, \
                patch.object(server, "OUTPUT_DIR", Path(directory)), \
                patch.object(server, "build_docx", side_effect=lambda _i, _t, p, **_k: p.write_bytes(b"word")), \
                patch.object(server, "convert_docx_to_pdf") as convert, \
                patch.object(server, "record_generated_document") as record:
            result = server.generate_proposal_download(self.context())
            self.assertTrue(result["filename"].endswith(".docx"))
            convert.assert_not_called()
            record.assert_called_once()

    def test_pdf_converts_the_same_template_and_edited_content(self):
        context = self.context()
        def convert(docx, pdf):
            self.assertEqual(docx.read_bytes(), b"word with template")
            pdf.write_bytes(b"%PDF-1.7\ncontent")
        with tempfile.TemporaryDirectory() as directory, \
                patch.object(server, "OUTPUT_DIR", Path(directory)), \
                patch.object(server, "build_docx", side_effect=lambda _i, _t, p, **_k: p.write_bytes(b"word with template")) as build, \
                patch.object(server, "convert_docx_to_pdf", side_effect=convert), \
                patch.object(server, "record_generated_document") as record:
            result = server.generate_proposal_download(context, "pdf")
            self.assertTrue(result["filename"].endswith(".pdf"))
            self.assertEqual(build.call_args.args[1], context["template_path"])
            for key in ("mini_box_order", "document_block_order", "mini_box_contents", "mini_box_alignments", "table_layout", "proposal_column_widths"):
                self.assertEqual(build.call_args.kwargs[key], context[key])
            self.assertEqual([p.suffix for p in Path(directory).iterdir()], [".pdf"])
            record.assert_called_once_with("1", Path(directory) / result["filename"])

    def test_conversion_failure_or_invalid_pdf_leaves_no_download_or_history(self):
        for failure in (True, False):
            def convert(_docx, pdf):
                pdf.write_bytes(b"not a PDF")
                if failure:
                    raise RuntimeError("converter unavailable")
            with self.subTest(failure=failure), tempfile.TemporaryDirectory() as directory, \
                    patch.object(server, "OUTPUT_DIR", Path(directory)), \
                    patch.object(server, "build_docx", side_effect=lambda _i, _t, p, **_k: p.write_bytes(b"word")), \
                    patch.object(server, "convert_docx_to_pdf", side_effect=convert), \
                    patch.object(server, "record_generated_document") as record:
                with self.assertRaisesRegex(ValueError, "PDF com o template"):
                    server.generate_proposal_download(self.context(), "pdf")
                self.assertEqual(list(Path(directory).iterdir()), [])
                record.assert_not_called()

    def test_invalid_format_is_rejected_before_building(self):
        with patch.object(server, "build_docx") as build:
            with self.assertRaises(ValueError):
                server.generate_proposal_download(self.context(), "html")
            build.assert_not_called()
