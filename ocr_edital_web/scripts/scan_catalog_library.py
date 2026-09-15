"""Scan retained catalog originals and index text with source locations for Block 7."""

import hashlib
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]


def scan_library():
    path = ROOT / "resources/catalog_library/references.json"
    manifest = json.loads(path.read_text(encoding="utf-8"))
    report = []
    for row in manifest["documents"]:
        name = unicodedata.normalize("NFKD", row["source"]).encode("ascii", "ignore").decode().lower()
        basename = Path(name).name
        if basename.startswith("edital ") or re.search(r"pedido(?: no)? edital", basename) or "analises de editais/" in name:
            row["kind"] = "requisito_edital"
        original = (ROOT / row["original"]).resolve()
        if not original.is_relative_to((ROOT / "data/catalog_repertoire/originals").resolve()):
            raise ValueError("Original outside catalog library")
        if hashlib.sha256(original.read_bytes()).hexdigest() != row["sha256"]:
            raise ValueError(f"Checksum mismatch: {row['source']}")
        locations = []
        try:
            if row["format"] == ".pdf":
                for index, page in enumerate(PdfReader(original).pages, 1):
                    locations.append({"page": index, "text": page.extract_text() or ""})
            elif row["format"] == ".docx":
                document = Document(original)
                for index, paragraph in enumerate(document.paragraphs, 1):
                    if paragraph.text.strip():
                        locations.append({"section": f"Parágrafo {index}", "text": paragraph.text})
                for index, table in enumerate(document.tables, 1):
                    for number, cells in enumerate(table.rows, 1):
                        locations.append({"section": f"Tabela {index}, linha {number}",
                                          "text": " | ".join(cell.text for cell in cells.cells)})
            else:
                locations = [{"section": "Relatório", "text": row["text"]}]
            row["text"] = "\n".join(location["text"] for location in locations)
            row["locations"] = locations
            row["scan_status"] = "text_extracted" if row["text"].strip() else "needs_ocr" if row["format"] == ".pdf" else "empty_document"
            row["empty_pages"] = [entry["page"] for entry in locations if "page" in entry and not entry["text"].strip()]
        except Exception as exc:
            row["scan_status"] = "extraction_error"
            row["scan_error"] = type(exc).__name__ + ": " + str(exc)[:300]
        report.append({key: row.get(key) for key in ("source", "kind", "scan_status", "empty_pages", "scan_error")})
    manifest["schema_version"] = 2
    inventory = next(row for row in manifest["documents"] if row["source"].endswith("validacao-modelos-cadeiras-por-categoria-v3.docx"))
    product_document = Document(ROOT / inventory["original"])
    categories = [row.cells[0].text for row in product_document.tables[0].rows[1:-1]]
    manifest["products"] = [
        {"name": row.cells[1].text, "category": category, "source_group": row.cells[2].text,
         "inventory_source": inventory["source"], "inventory_sha256": inventory["sha256"],
         "status": "documented_reference"}
        for category, table in zip(categories, product_document.tables[1:]) for row in table.rows[1:]
    ]
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    output = ROOT / "reports/catalog-library/scan.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    summary = {"documents": len(report), "documented_products": len(manifest["products"]), "states": dict(Counter(row["scan_status"] for row in report)),
               "searchable": sum(bool(row["text"].strip()) and row["kind"] != "requisito_edital" for row in manifest["documents"]),
               "documents_detail": report}
    output.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: value for key, value in summary.items() if key != "documents_detail"}))


if __name__ == "__main__":
    scan_library()
