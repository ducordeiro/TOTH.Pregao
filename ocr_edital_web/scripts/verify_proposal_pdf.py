"""Exercise the real Word converter without writing to the production database."""
import json
import sys
import time
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pypdfium2
from pypdf import PdfReader
from docx import Document
import server

out = server.ROOT / "reports" / "proposal-pdf"
out.mkdir(parents=True, exist_ok=True)
template = out / "template-test.docx"
doc = Document()
doc.sections[0].header.paragraphs[0].text = "CABECALHO DO TEMPLATE"
doc.sections[0].footer.paragraphs[0].text = "RODAPE DO TEMPLATE"
doc.add_paragraph("{PROPOSTA DE TESTE}")
doc.save(template)
structure = server.inspect_docx_structure(template)
ids = [node["id"] for node in structure["nodes"] if node["type"] == "MINI_BOX"]
context = {
    "items": [{"item": "2", "quantidade": "32", "unidade": "UND", "descricao": "Cadeira de teste com assento em espuma laminada, estrutura em aco, encosto ajustavel e acabamento pintado em epoxi.",
               "marca": "rftgyh", "valor_unitario": "R$ 234,00", "valor_total": "R$ 7.488,00"}],
    "template_path": template, "source_name": "teste_pdf", "responsible_id": "1",
    "responsible": None, "commercial_terms": {}, "mini_box_order": ids,
    "document_block_order": [*ids, structure["generated_table_block"]["id"]],
    "mini_box_alignments": {ids[0]: "center"}, "mini_box_contents": {ids[0]: "PROPOSTA EDITADA"},
    "extra_column": None, "table_layout": None, "proposal_column_widths": {
        "item": 6, "quantidade": 6, "unidade": 6, "descricao": 55,
        "marca": 8, "valor_unitario": 9.5, "valor_total": 9.5,
    },
}
started = time.perf_counter()
with patch.object(server, "OUTPUT_DIR", out), patch.object(server, "record_generated_document"):
    result = server.generate_proposal_download(context, "pdf")
with pypdfium2.PdfDocument(out / result["filename"]) as pdf:
    text = "\n".join(page.extract_text() for page in PdfReader(out / result["filename"]).pages)
    for expected in ("CABECALHO DO TEMPLATE", "RODAPE DO TEMPLATE", "PROPOSTA EDITADA", "Cadeira de teste"):
        assert expected in text, (expected, text)
    assert "{" not in text and "}" not in text
    assert "7.488,00" in text, text
    pdf[0].render(scale=1).to_pil().save(out / "preview.png")
    print(json.dumps({"file": result["filename"], "pages": len(pdf), "seconds": round(time.perf_counter() - started, 2), "checks": "header, footer, table, edited mini-box and removed braces"}))
