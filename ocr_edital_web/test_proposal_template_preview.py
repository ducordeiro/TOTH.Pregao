import base64
import io
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from proposal_template_preview import template_page_parts_preview


class TemplatePagePartsTests(unittest.TestCase):
    def test_original_header_logo_styles_and_breaks_are_preserved(self):
        source = Path(__file__).parent / "templates/modelo_proposta_goldflex.docx"
        original_bytes = source.read_bytes()
        result = template_page_parts_preview(source)
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(io.BytesIO(base64.b64decode(result["docx_base64"]))) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != "word/document.xml":
                    self.assertEqual(before.read(name), after.read(name), name)
            self.assertIn(b"drawing", after.read("word/header1.xml"))
            self.assertIn(b'<w:jc w:val="center"', after.read("word/header1.xml"))
            self.assertGreaterEqual(after.read("word/header1.xml").count(b"<w:br"), 3)
            self.assertEqual(after.read("word/document.xml").count(b'w:type="page"'), 2)
        self.assertEqual(source.read_bytes(), original_bytes)
        self.assertAlmostEqual(result["content_width_pt"], 451.3, places=1)

    def test_inherited_headers_and_first_even_variants_remain_available(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "template.docx"
            document = Document()
            first = document.sections[0]
            first.header.paragraphs[0].text = "DEFAULT HEADER"
            first.first_page_header.paragraphs[0].text = "FIRST HEADER"
            first.even_page_header.paragraphs[0].text = "EVEN HEADER"
            first.footer.paragraphs[0].text = "DEFAULT FOOTER"
            first.first_page_footer.paragraphs[0].text = "FIRST FOOTER"
            first.even_page_footer.paragraphs[0].text = "EVEN FOOTER"
            document.settings.odd_and_even_pages_header_footer = True
            document.add_section(WD_SECTION.NEW_PAGE).different_first_page_header_footer = True
            document.save(path)
            result = template_page_parts_preview(path)
            preview = Document(io.BytesIO(base64.b64decode(result["docx_base64"])))
            self.assertEqual(len(preview.sections), 1)
            section = preview.sections[0]
            self.assertEqual(section.header.paragraphs[0].text, "DEFAULT HEADER")
            self.assertEqual(section.first_page_header.paragraphs[0].text, "FIRST HEADER")
            self.assertEqual(section.even_page_header.paragraphs[0].text, "EVEN HEADER")
            self.assertTrue(section.different_first_page_header_footer)
            self.assertFalse(result["first_header_blank"])

    def test_template_change_invalidates_cached_preview(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "template.docx"
            document = Document()
            document.sections[0].header.paragraphs[0].text = "First"
            document.save(path)
            before = template_page_parts_preview(path)
            document.sections[0].header.paragraphs[0].text = "Updated header"
            document.save(path)
            after = template_page_parts_preview(path)
            self.assertNotEqual(before["docx_base64"], after["docx_base64"])
            self.assertEqual(template_page_parts_preview(path), after)


if __name__ == "__main__":
    unittest.main()
