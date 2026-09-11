import tempfile
import unittest
from pathlib import Path
from docx import Document
import server


class ProposalTableEditorTests(unittest.TestCase):
    def setUp(self):
        self.items = [dict(item='1', quantidade='2', unidade='CAIXA', descricao='Produto editado',
                           marca='Marca', valor_unitario='R$ 10,00', valor_total='R$ 20,00')]
        self.layout = {'columns': [
            {'key': 'descricao', 'label': 'Produto'}, {'key': 'unidade', 'label': 'Embalagem'},
            {'key': 'custom_a', 'label': 'Modelo'}, {'key': 'custom_b', 'label': 'Garantia'}],
            'custom_values': {'custom_a': ['X-1'], 'custom_b': ['12 meses\nAssistencia']}}

    def test_docx_reflects_saved_columns_and_values_preserving_template(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template, output = root / 'template.docx', root / 'output.docx'
            doc = Document()
            doc.sections[0].header.paragraphs[0].text = 'Cabecalho do modelo'
            doc.add_paragraph('Texto fixo do modelo')
            doc.save(template)
            server.build_docx(self.items, template, output, table_layout=self.layout)
            generated = Document(output)
            self.assertEqual([c.text for c in generated.tables[0].rows[0].cells], ['Produto', 'Embalagem', 'Modelo', 'Garantia'])
            self.assertEqual([c.text for c in generated.tables[0].rows[1].cells], ['Produto editado', 'CAIXA', 'X-1', '12 meses\nAssistencia'])
            self.assertEqual(generated.sections[0].header.paragraphs[0].text, 'Cabecalho do modelo')
            self.assertIn('Texto fixo do modelo', [p.text for p in generated.paragraphs])

    def test_rejects_invalid_layouts(self):
        for layout in (
            {'columns': [], 'custom_values': {}},
            self.layout | {'columns': self.layout['columns'] * 2},
            self.layout | {'custom_values': {'custom_a': [], 'custom_b': ['x']}},
            {'columns': [{'key': 'unknown', 'label': 'X'}]},
            self.layout | {'custom_values': {'custom_a': ['x\x00'], 'custom_b': ['x']}},
        ):
            with self.subTest(layout=layout), self.assertRaises(ValueError):
                server.normalize_proposal_table_layout(self.items, layout)

    def test_widths_and_cache_follow_layout(self):
        a = server.proposal_preview_fingerprint({'items': self.items, 'table_layout': self.layout})
        b = server.proposal_preview_fingerprint({'items': self.items, 'table_layout': None})
        self.assertNotEqual(a, b)
        widths = server.normalize_proposal_column_widths(self.items, {}, table_layout=self.layout)
        self.assertEqual(set(widths), {'descricao', 'unidade', 'custom_a', 'custom_b'})
        self.assertAlmostEqual(sum(widths.values()), 100)


if __name__ == '__main__':
    unittest.main()
